#!/usr/bin/env python3
"""Validate the consolidated cbenef-finder base against the current rule engine."""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


DEFAULT_INPUT = Path(r"C:/Users/thatf/Downloads/cbenef_finder_base_validacao_enxuta.xlsx")
DEFAULT_OUTPUT_DIR = Path("outputs/cbenef-validation")
DEFAULT_SHEET = "Base_Validacao"
DEFAULT_BUILDER = Path("scripts/build_cbenef_validation_workbook.mjs")

REQUESTED_COLUMNS = [
    "Código",
    "Descrição",
    "NCM",
    "CST_ICMS<S>",
    "% ICMS<S>",
    "Trib.",
    "CFOP",
    "cBenef",
    "CEST",
    "Dep.Nome",
    "Tipo produto",
]

OUTPUT_COLUMNS = [
    "Código",
    "Descrição",
    "NCM_esperado",
    "NCM_retornado",
    "CST_esperado",
    "CST_retornado",
    "%ICMS_esperado",
    "%ICMS_retornado",
    "TRIB_esperado",
    "TRIB_retornado",
    "CFOP_esperado",
    "CFOP_retornado",
    "cBenef_esperado",
    "cBenef_retornado",
    "status_validacao",
    "motivo_divergencia",
    "CEST",
    "Dep.Nome",
    "Tipo produto",
]

NOISE_TOKENS = {
    "promissao",
    "aurora",
    "tirolez",
    "frizzo",
    "criolo",
    "santo",
    "antonio",
    "parmalat",
    "nestle",
    "danone",
    "elegante",
    "italac",
    "piracanjuba",
    "vigor",
    "presidente",
    "pote",
    "pct",
    "pct.",
    "pcts",
    "un",
    "und",
    "unid",
    "cx",
    "cxa",
    "caixa",
    "kg",
    "gr",
    "g",
    "ml",
    "lt",
    "lts",
    "litro",
    "litros",
    "100",
    "150",
    "200",
    "250",
    "300",
    "400",
    "500",
    "600",
    "750",
    "1000",
    "de",
    "do",
    "da",
    "com",
    "sem",
    "para",
    "em",
    "no",
    "na",
    "os",
    "as",
    "ao",
    "pela",
    "pelo",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Validate the consolidated cbenef-finder base against the current rule engine.",
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help="Path to the consolidated validation workbook.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="Directory where the JSON and XLSX validation artifacts will be written.",
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help="Worksheet name that contains the consolidated base.",
    )
    parser.add_argument(
        "--builder",
        type=Path,
        default=DEFAULT_BUILDER,
        help="Path to the Node workbook builder helper.",
    )
    parser.add_argument(
        "--max-rows",
        type=int,
        default=None,
        help="Optional row limit for a quick smoke test.",
    )
    return parser.parse_args()


def parse_env_file(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}

    result: dict[str, str] = {}
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        result[key] = value
    return result


def resolve_supabase_config() -> tuple[str, str]:
    env = {**parse_env_file(Path(".env")), **os.environ}
    url = env.get("VITE_SUPABASE_URL", "").strip()
    key = env.get("VITE_SUPABASE_PUBLISHABLE_KEY", "").strip()
    if not url or not key:
        raise RuntimeError(
            "Missing Supabase configuration. Set VITE_SUPABASE_URL and VITE_SUPABASE_PUBLISHABLE_KEY.",
        )
    return url.rstrip("/"), key


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "g")
    return str(value).strip()


def stringify_number(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    try:
        number = Decimal(text.replace(",", "."))
    except InvalidOperation:
        return text
    normalized = number.normalize()
    rendered = format(normalized, "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered or "0"


def stringify_code(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    try:
        number = float(text.replace(",", "."))
    except ValueError:
        return text
    if number.is_integer():
        return str(int(number))
    return stringify_number(text)


def remove_accents(text: Any) -> str:
    raw = clean_text(text)
    if not raw:
        return ""
    normalized = unicodedata.normalize("NFD", raw)
    return "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")


def collapse_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def normalize_description_key(value: Any) -> str:
    return collapse_spaces(remove_accents(value).upper())


def normalize_description_request(value: Any) -> str:
    return clean_text(value)


def normalize_digits(value: Any) -> str:
    return re.sub(r"\D+", "", clean_text(value))


def normalize_ncm(value: Any) -> str:
    return normalize_digits(value)


def normalize_cst(value: Any) -> str:
    digits = normalize_digits(value)
    if not digits:
        return ""
    return digits.zfill(3)


def normalize_cfop(value: Any) -> str:
    digits = normalize_digits(value)
    if not digits:
        return ""
    return digits.zfill(4)


def normalize_trib(value: Any) -> str:
    return collapse_spaces(remove_accents(value).upper())


def normalize_cbenef(value: Any) -> str:
    text = collapse_spaces(remove_accents(value).upper())
    return text.replace(" ", "")


def parse_decimal(value: Any) -> Decimal | None:
    text = clean_text(value)
    if not text:
        return None
    text = text.replace("%", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})*(?:,\d+)?", text):
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def normalize_percent(value: Any) -> str:
    parsed = parse_decimal(value)
    if parsed is None:
        return ""
    rendered = format(parsed.normalize(), "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered or "0"


def compare_percent(expected: Any, returned: Any) -> bool:
    left = parse_decimal(expected)
    right = parse_decimal(returned)
    if left is None or right is None:
        return normalize_percent(expected) == normalize_percent(returned)
    return left == right


def load_workbook_rows(path: Path, sheet_name: str, max_rows: int | None) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise KeyError(f"Sheet {sheet_name!r} not found in {path}")

    sheet = workbook[sheet_name]
    rows: list[dict[str, Any]] = []
    header_row = None
    column_index: dict[str, int] = {}

    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index == 1:
            header_row = [clean_text(cell) for cell in row]
            column_index = {name: pos for pos, name in enumerate(header_row)}
            missing = [name for name in REQUESTED_COLUMNS if name not in column_index]
            if missing:
                raise KeyError(
                    f"Missing required columns in {sheet_name!r}: {', '.join(missing)}",
                )
            continue

        if max_rows is not None and len(rows) >= max_rows:
            break

        record = {column: row[column_index[column]] if column_index[column] < len(row) else None for column in REQUESTED_COLUMNS}
        record["_source_row"] = index
        rows.append(record)

    if header_row is None:
        raise RuntimeError(f"Worksheet {sheet_name!r} is empty.")
    return rows


def build_validation_key(row: dict[str, Any]) -> str:
    return f"{normalize_ncm(row['NCM'])}|{normalize_description_key(row['Descrição'])}"


def fetch_json_page(base_url: str, table: str, params: dict[str, Any], headers: dict[str, str]) -> list[dict[str, Any]]:
    query = urlencode({key: value for key, value in params.items() if value is not None})
    url = f"{base_url}/rest/v1/{table}"
    if query:
        url = f"{url}?{query}"
    request = Request(url, headers=headers, method="GET")
    with urlopen(request, timeout=120) as response:
        payload = response.read().decode("utf-8")
    data = json.loads(payload)
    if not isinstance(data, list):
        raise RuntimeError(f"Unexpected response for {table}: {data!r}")
    return data


def fetch_all_table(
    base_url: str,
    table: str,
    headers: dict[str, str],
    *,
    select: str = "*",
    filters: dict[str, Any] | None = None,
    page_size: int = 1000,
) -> list[dict[str, Any]]:
    params: dict[str, Any] = {"select": select, "limit": page_size}
    if filters:
        params.update(filters)

    rows: list[dict[str, Any]] = []
    offset = 0
    while True:
        params["offset"] = offset
        page = fetch_json_page(base_url, table, params, headers)
        rows.extend(page)
        if len(page) < page_size:
            break
        offset += page_size
    return rows


@dataclass(slots=True)
class TaxonomyEntry:
    macro_code: str
    sub_code: str
    keyword: str
    match_type: str
    weight: float


def build_taxonomy(
    active_groups: list[dict[str, Any]],
    keywords: list[dict[str, Any]],
) -> list[TaxonomyEntry]:
    id_to_code = {clean_text(row["id"]): clean_text(row["code"]) for row in active_groups}
    id_to_parent = {clean_text(row["id"]): clean_text(row.get("parent_id")) for row in active_groups}

    entries: list[TaxonomyEntry] = []
    for row in keywords:
        group_id = clean_text(row.get("group_id"))
        group_code = id_to_code.get(group_id)
        if not group_code:
            continue

        parent_id = id_to_parent.get(group_id)
        if parent_id and id_to_code.get(parent_id):
            macro_code = id_to_code[parent_id]
            sub_code = group_code
        else:
            macro_code = group_code
            sub_code = ""

        weight = row.get("weight", 0) or 0
        try:
            weight_value = float(weight)
        except (TypeError, ValueError):
            weight_value = 0.0

        entries.append(
            TaxonomyEntry(
                macro_code=macro_code,
                sub_code=sub_code,
                keyword=normalize_description_key(row.get("keyword")).lower(),
                match_type=clean_text(row.get("match_type")),
                weight=weight_value,
            )
        )
    return entries


def infer_group_from_taxonomy(
    strong_tokens: list[str],
    normalized_desc: str,
    taxonomy: list[TaxonomyEntry],
) -> dict[str, str]:
    scores: dict[str, dict[str, Any]] = {}

    for entry in taxonomy:
        if entry.match_type != "include":
            continue

        parts = [part for part in entry.keyword.split(" ") if part]
        hit = False
        if len(parts) > 1:
            if all(part in strong_tokens for part in parts) or entry.keyword in normalized_desc:
                hit = True
        elif entry.keyword in strong_tokens or entry.keyword in normalized_desc:
            hit = True

        if not hit:
            continue

        key = entry.sub_code or entry.macro_code
        if key not in scores:
            scores[key] = {
                "macro": entry.macro_code,
                "sub": entry.sub_code,
                "score": 0.0,
            }
        scores[key]["score"] += 10 * entry.weight

    best_macro = ""
    best_sub = ""
    best_score = 0.0
    for candidate in scores.values():
        if candidate["score"] > best_score:
            best_macro = candidate["macro"]
            best_sub = candidate["sub"]
            best_score = candidate["score"]

    return {
        "inferred_macro_group": best_macro,
        "inferred_subgroup": best_sub,
    }


def infer_product_classification(strong_tokens: list[str]) -> dict[str, str]:
    joined = " ".join(strong_tokens)
    family = ""
    product_type = ""
    presentation = ""

    cheese_types: dict[str, list[str]] = {
        "mussarela": ["mussarela", "mucarela", "mucarel", "mussarel"],
        "prato": ["queijo prato"],
        "provolone": ["provolone"],
        "coalho": ["coalho", "queijo coalho"],
        "minas_frescal": ["minas frescal", "minas", "frescal"],
        "queijo_ralado": ["queijo ralado"],
    }
    for dtype, keywords in cheese_types.items():
        if any(
            keyword in joined
            or all(part in strong_tokens for part in keyword.split(" "))
            for keyword in keywords
        ):
            family = "queijos"
            product_type = dtype
            break

    if not family:
        dairy_types: dict[str, list[str]] = {
            "manteiga": ["manteiga"],
            "requeijao": ["requeijao"],
            "ricota": ["ricota"],
            "iogurte": ["iogurte", "yogurte"],
            "leite": ["leite"],
            "creme_de_leite": ["creme leite"],
            "nata": ["nata"],
            "leite_condensado": ["leite condensado"],
        }
        for dtype, keywords in dairy_types.items():
            if any(
                keyword in joined
                or all(part in strong_tokens for part in keyword.split(" "))
                for keyword in keywords
            ):
                family = "laticinios"
                product_type = dtype
                break

    if not family:
        meat_kw: dict[str, dict[str, list[str]]] = {
            "bovino": {
                "family": "carnes_bovinas",
                "kws": [
                    "alcatra",
                    "picanha",
                    "patinho",
                    "file mignon",
                    "acem",
                    "coxao",
                    "maminha",
                    "fraldinha",
                    "lagarto",
                    "musculo",
                    "costela bovina",
                    "charque",
                    "carne seca",
                    "carne sol",
                ],
            },
            "suino": {
                "family": "carnes_suinas",
                "kws": [
                    "pernil",
                    "lombo suino",
                    "bisteca suina",
                    "costela suina",
                    "bacon",
                    "toucinho",
                    "panceta",
                ],
            },
            "frango": {
                "family": "aves",
                "kws": [
                    "frango",
                    "coxa",
                    "sobrecoxa",
                    "asa",
                    "peito frango",
                    "sassami",
                    "moela",
                    "coracao frango",
                ],
            },
            "peixe": {
                "family": "pescados",
                "kws": [
                    "tilapia",
                    "merluza",
                    "sardinha",
                    "salmao",
                    "bacalhau",
                    "corvina",
                    "peixe",
                    "file peixe",
                ],
            },
        }
        for dtype, cfg in meat_kw.items():
            if any(keyword in joined for keyword in cfg["kws"]):
                family = cfg["family"]
                product_type = dtype
                break

    presentations: dict[str, list[str]] = {
        "fatiada": ["fatiada", "fatiado", "fatia", "fatiados"],
        "pedaco": ["pedaco", "peca", "inteiro", "inteira", "bloco"],
        "ralado": ["ralado", "ralada"],
        "pote": ["pote"],
        "tablete": ["tablete", "tabletes", "barra"],
        "banda": ["banda"],
    }
    for ptype, keywords in presentations.items():
        if any(keyword in strong_tokens for keyword in keywords):
            presentation = ptype
            break

    if not presentation and product_type:
        presentation = "padrao"
    if not family and not product_type:
        family = "nao_identificado"
        product_type = "nao_identificado"

    return {
        "product_family": family,
        "product_type": product_type,
        "presentation_type": presentation,
    }


def score_rule(
    rule: dict[str, Any],
    strong_tokens: list[str],
    normalized_desc: str,
    inferred_type: str,
    inferred_family: str,
    inferred_presentation: str,
    inferred_macro: str,
    inferred_sub: str,
    informed_group: str | None,
) -> dict[str, Any]:
    score = 0.0
    exclude_hit = False
    include_hits: list[str] = []

    kw_exclude = [remove_accents(value).lower() for value in rule.get("keyword_exclude") or []]
    for token in kw_exclude:
        if token in strong_tokens or token in normalized_desc:
            exclude_hit = True

    kw_include = [remove_accents(value).lower() for value in rule.get("keyword_include") or []]
    for token in kw_include:
        parts = [part for part in token.split(" ") if part]
        if len(parts) > 1:
            if all(part in strong_tokens or part in normalized_desc for part in parts):
                score += 10
                include_hits.append(token)
        elif token in strong_tokens or token in normalized_desc:
            score += 10
            include_hits.append(token)

    for token in [remove_accents(value).lower() for value in rule.get("keywords") or []]:
        if token in strong_tokens or token in normalized_desc:
            score += 2

    for pattern in [remove_accents(value).lower() for value in rule.get("description_patterns") or []]:
        if pattern in normalized_desc:
            score += 8

    product_type_match = False
    rule_product_type = remove_accents(rule.get("product_type")).lower()
    if rule_product_type and inferred_type:
        if rule_product_type == inferred_type:
            score += 15
            product_type_match = True

    rule_family = remove_accents(rule.get("product_family")).lower()
    if rule_family and inferred_family and rule_family == inferred_family:
        score += 5

    presentation_match = False
    rule_presentation = remove_accents(rule.get("presentation_type")).lower()
    if rule_presentation and inferred_presentation:
        if rule_presentation == inferred_presentation:
            score += 12
            presentation_match = True
        else:
            score -= 5

    group_match = False
    rule_macro = remove_accents(rule.get("macro_group")).lower()
    if rule_macro:
        if rule_macro == inferred_macro:
            score += 8
            group_match = True
        if informed_group and remove_accents(informed_group).lower() == rule_macro:
            score += 5
            group_match = True

    rule_subgroup = remove_accents(rule.get("subgroup")).lower()
    if rule_subgroup and inferred_sub and rule_subgroup == inferred_sub:
        score += 6

    if clean_text(rule.get("data_origin")) == "imported":
        score += 3

    try:
        priority = int(rule.get("priority", 0) or 0)
    except (TypeError, ValueError):
        priority = 0
    score += min(priority, 20)

    return {
        "rule": rule,
        "score": score,
        "excludeHit": exclude_hit,
        "includeHits": include_hits,
        "productTypeMatch": product_type_match,
        "groupMatch": group_match,
        "presentationMatch": presentation_match,
    }


def build_low_confidence_response(
    ncm: str,
    informed_cst: str | None,
    normalized_description: dict[str, Any],
    classification: dict[str, str],
    group_inference: dict[str, str],
    informed_group: str | None,
) -> dict[str, Any]:
    return {
        "cbenef_code": "",
        "informed_cst_icms": informed_cst or "",
        "suggested_cst_icms": "",
        "final_cst_icms": "",
        "cst_source": "none",
        "confidence_score": 0,
        "confidence_level": "low",
        "matched_rule_id": None,
        "application_context": "Operacao interna - Estado de Sao Paulo",
        "legal_basis_name": "",
        "legal_basis_summary": "",
        "legal_basis_url": None,
        "rule_version": None,
        "last_updated_at": datetime.now(timezone.utc).isoformat(),
        "input_ncm": ncm,
        "matched_ncm": "",
        "explanation": "",
        "matched_by_ncm_exact": False,
        "matched_by_ncm_prefix": False,
        "keyword_match_count": 0,
        "used_informed_cst": False,
        "auto_suggested_cst": False,
        "data_origin": "",
        "normalized_description": normalized_description["normalized_description"],
        "matched_keywords": [],
        "excluded_keywords_hit": [],
        "inferred_macro_group": group_inference["inferred_macro_group"],
        "inferred_subgroup": group_inference["inferred_subgroup"],
        "informed_group": informed_group or "",
        "group_consistency_status": "",
        "product_family": classification["product_family"],
        "product_type": classification["product_type"],
        "presentation_type": classification["presentation_type"],
        "output_st_applicable": None,
        "output_cfop": "",
        "output_trib_code": "",
        "output_icms_rate": None,
        "decision_reason": "",
    }


def build_engine_tables(
    rules: list[dict[str, Any]],
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    rules_by_exact: dict[str, list[dict[str, Any]]] = defaultdict(list)
    rules_by_prefix: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for rule in rules:
        ncm = normalize_ncm(rule.get("ncm"))
        if not ncm:
            continue
        rule["_ncm_norm"] = ncm
        rule["_cst_icms_norm"] = normalize_cst(rule.get("cst_icms"))
        rule["_suggested_cst_icms_norm"] = normalize_cst(rule.get("suggested_cst_icms"))
        rule["_output_cst_icms_norm"] = normalize_cst(rule.get("output_cst_icms"))
        rule["_output_cfop_norm"] = normalize_cfop(rule.get("output_cfop"))
        rule["_output_trib_code_norm"] = normalize_trib(rule.get("output_trib_code"))
        rules_by_exact[ncm].append(rule)
        for prefix_len in (2, 4, 6, 8):
            if len(ncm) >= prefix_len:
                rules_by_prefix[ncm[:prefix_len]].append(rule)

    return rules_by_exact, rules_by_prefix


def classify_row(
    row: dict[str, Any],
    rules_by_exact: dict[str, list[dict[str, Any]]],
    rules_by_prefix: dict[str, list[dict[str, Any]]],
    taxonomy: list[TaxonomyEntry],
) -> dict[str, Any]:
    raw_description = normalize_description_request(row["Descrição"])
    ncm = normalize_ncm(row["NCM"])
    informed_cst = normalize_cst(row["CST_ICMS<S>"])

    normalized_description = collapse_spaces(remove_accents(raw_description).lower())
    normalized_tokens = [token for token in normalized_description.split(" ") if token]
    strong_tokens: list[str] = []
    for token in normalized_tokens:
        if token in NOISE_TOKENS or re.fullmatch(r"\d+", token):
            continue
        strong_tokens.append(token)

    classification = infer_product_classification(strong_tokens)
    group_inference = infer_group_from_taxonomy(strong_tokens, normalized_description, taxonomy)

    desc_insufficient = len(strong_tokens) < 2 or len([token for token in strong_tokens if token not in {
        "produto",
        "alimenticio",
        "alimento",
        "item",
        "mercadoria",
        "material",
        "laticinio",
        "laticinios",
        "carne",
        "peixe",
        "queijo",
        "mercearia",
    }]) < 1

    exact_rules = rules_by_exact.get(ncm, [])
    matched_by_ncm_exact = bool(exact_rules)
    matched_by_ncm_prefix = False
    candidate_rules = exact_rules

    if not candidate_rules:
        for prefix in (ncm[:6], ncm[:4], ncm[:2]):
            candidate_rules = rules_by_prefix.get(prefix, [])
            if candidate_rules:
                matched_by_ncm_prefix = True
                break

    if not candidate_rules:
        response = build_low_confidence_response(ncm, informed_cst, {
            "normalized_description": normalized_description,
        }, classification, group_inference, None)
        response["matched_by_ncm_exact"] = matched_by_ncm_exact
        response["matched_by_ncm_prefix"] = matched_by_ncm_prefix
        return response

    scored = [
        score_rule(
            rule,
            strong_tokens,
            normalized_description,
            classification["product_type"],
            classification["product_family"],
            classification["presentation_type"],
            group_inference["inferred_macro_group"],
            group_inference["inferred_subgroup"],
            None,
        )
        for rule in candidate_rules
    ]

    eligible = [entry for entry in scored if not entry["excludeHit"]]
    excluded_rules = [entry for entry in scored if entry["excludeHit"]]

    if not eligible:
        response = build_low_confidence_response(ncm, informed_cst, {
            "normalized_description": normalized_description,
        }, classification, group_inference, None)
        response["matched_by_ncm_exact"] = matched_by_ncm_exact
        response["matched_by_ncm_prefix"] = matched_by_ncm_prefix
        return response

    eligible.sort(
        key=lambda entry: (
            0 if entry["presentationMatch"] else 1,
            0 if entry["productTypeMatch"] else 1,
            0 if entry["groupMatch"] else 1,
            -len(entry["includeHits"]),
            -entry["score"],
            -int(entry["rule"].get("priority", 0) or 0),
        )
    )

    best = eligible[0]
    best_rule = best["rule"]

    rule_cst = (
        normalize_cst(best_rule.get("output_cst_icms"))
        or normalize_cst(best_rule.get("suggested_cst_icms"))
        or normalize_cst(best_rule.get("cst_icms"))
        or ""
    )

    cst_source = "sugerido"
    final_cst = ""
    used_informed_cst = False
    auto_suggested_cst = False

    if informed_cst:
        if informed_cst == rule_cst or informed_cst == normalize_cst(best_rule.get("cst_icms")):
            cst_source = "informado"
            final_cst = informed_cst
            used_informed_cst = True
        else:
            cst_source = "ajustado"
            final_cst = rule_cst
    else:
        final_cst = rule_cst
        auto_suggested_cst = True

    response = {
        "cbenef_code": clean_text(best_rule.get("cbenef_code")),
        "informed_cst_icms": informed_cst or "",
        "suggested_cst_icms": normalize_cst(best_rule.get("suggested_cst_icms"))
        or normalize_cst(best_rule.get("cst_icms"))
        or "",
        "final_cst_icms": final_cst,
        "cst_source": cst_source,
        "confidence_score": None,
        "confidence_level": "",
        "matched_rule_id": clean_text(best_rule.get("id")),
        "application_context": clean_text(best_rule.get("application_context"))
        or "Operacao interna - Estado de Sao Paulo",
        "legal_basis_name": clean_text(best_rule.get("legal_basis_name"))
        or clean_text(best_rule.get("legal_basis"))
        or "",
        "legal_basis_summary": clean_text(best_rule.get("legal_basis_summary")),
        "legal_basis_url": best_rule.get("legal_basis_url") or best_rule.get("legal_url") or None,
        "rule_version": None,
        "last_updated_at": clean_text(best_rule.get("updated_at")) or clean_text(best_rule.get("created_at")),
        "input_ncm": ncm,
        "matched_ncm": normalize_ncm(best_rule.get("ncm")),
        "explanation": "",
        "matched_by_ncm_exact": matched_by_ncm_exact,
        "matched_by_ncm_prefix": matched_by_ncm_prefix,
        "keyword_match_count": len(best["includeHits"]),
        "used_informed_cst": used_informed_cst,
        "auto_suggested_cst": auto_suggested_cst,
        "data_origin": clean_text(best_rule.get("data_origin")),
        "normalized_description": normalized_description,
        "matched_keywords": best["includeHits"],
        "excluded_keywords_hit": [
            token
            for entry in excluded_rules
            for token in (entry["rule"].get("keyword_exclude") or [])
            if remove_accents(token).lower() in normalized_description
            or remove_accents(token).lower() in strong_tokens
        ],
        "inferred_macro_group": group_inference["inferred_macro_group"],
        "inferred_subgroup": group_inference["inferred_subgroup"],
        "informed_group": "",
        "group_consistency_status": "nao_identificado",
        "product_family": clean_text(best_rule.get("product_family")) or classification["product_family"],
        "product_type": clean_text(best_rule.get("product_type")) or classification["product_type"],
        "presentation_type": clean_text(best_rule.get("presentation_type")) or classification["presentation_type"],
        "output_st_applicable": best_rule.get("output_st_applicable"),
        "output_cfop": normalize_cfop(best_rule.get("output_cfop")),
        "output_trib_code": normalize_trib(best_rule.get("output_trib_code")),
        "output_icms_rate": best_rule.get("output_icms_rate"),
        "decision_reason": clean_text(best_rule.get("decision_reason")),
    }

    if desc_insufficient:
        response["confidence_level"] = "low"
    else:
        response["confidence_level"] = "medium"

    return response


def build_conflict_map(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[build_validation_key(row)].append(row)

    conflicts: dict[str, dict[str, Any]] = {}
    for key, group in grouped.items():
        if len(group) <= 1:
            continue

        varying: dict[str, list[str]] = {}
        for field, normalizer in (
            ("NCM", normalize_ncm),
            ("CST_ICMS<S>", normalize_cst),
            ("% ICMS<S>", normalize_percent),
            ("Trib.", normalize_trib),
            ("CFOP", normalize_cfop),
            ("cBenef", normalize_cbenef),
        ):
            values = sorted({normalizer(item[field]) for item in group if normalizer(item[field])})
            if len(values) > 1:
                varying[field] = values

        if varying:
            reasons = []
            for field, values in varying.items():
                reasons.append(f"{field} divergente ({', '.join(values)})")
            conflicts[key] = {
                "rows": [item["_source_row"] for item in group],
                "reason": "Conflito interno na base: " + "; ".join(reasons),
            }

    return conflicts


def compare_row(
    row: dict[str, Any],
    response: dict[str, Any],
    conflict_reason: str | None,
) -> tuple[str, str]:
    if conflict_reason:
        return "INCONCLUSIVO", conflict_reason

    mismatches: list[str] = []

    expected_ncm = normalize_ncm(row["NCM"])
    returned_ncm = normalize_ncm(response.get("matched_ncm"))
    if expected_ncm != returned_ncm:
        mismatches.append(f"NCM esperado {expected_ncm} vs retornado {returned_ncm or '[vazio]'}")

    expected_cst = normalize_cst(row["CST_ICMS<S>"])
    returned_cst = normalize_cst(response.get("final_cst_icms"))
    if expected_cst != returned_cst:
        mismatches.append(f"CST esperado {expected_cst} vs retornado {returned_cst or '[vazio]'}")

    if not compare_percent(row["% ICMS<S>"], response.get("output_icms_rate")):
        mismatches.append(
            f"%ICMS esperado {normalize_percent(row['% ICMS<S>'])} vs retornado {normalize_percent(response.get('output_icms_rate')) or '[vazio]'}",
        )

    expected_trib = normalize_trib(row["Trib."])
    returned_trib = normalize_trib(response.get("output_trib_code"))
    if expected_trib != returned_trib:
        mismatches.append(f"TRIB esperado {expected_trib} vs retornado {returned_trib or '[vazio]'}")

    expected_cfop = normalize_cfop(row["CFOP"])
    returned_cfop = normalize_cfop(response.get("output_cfop"))
    if expected_cfop != returned_cfop:
        mismatches.append(f"CFOP esperado {expected_cfop} vs retornado {returned_cfop or '[vazio]'}")

    expected_cbenef = normalize_cbenef(row["cBenef"])
    if expected_cbenef:
        returned_cbenef = normalize_cbenef(response.get("cbenef_code"))
        if expected_cbenef != returned_cbenef:
            mismatches.append(
                f"cBenef esperado {expected_cbenef} vs retornado {returned_cbenef or '[vazio]'}",
            )

    if mismatches:
        return "DIVERGIU", "; ".join(mismatches)
    return "BATEU", ""


def main() -> int:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    rows = load_workbook_rows(args.input, args.sheet, args.max_rows)
    conflict_map = build_conflict_map(rows)
    conflict_row_numbers = {
        row_number
        for conflict in conflict_map.values()
        for row_number in conflict["rows"]
    }

    supabase_cache = args.output_dir / "supabase_tables.json"
    fetcher = Path("scripts/fetch_cbenef_supabase.mjs").resolve()
    subprocess.run(
        [
            "node",
            str(fetcher),
            "--output-json",
            str(supabase_cache),
        ],
        check=True,
    )

    supabase_payload = json.loads(supabase_cache.read_text(encoding="utf-8"))
    active_groups = supabase_payload["classification_groups"]
    group_keywords = supabase_payload["classification_group_keywords"]
    rules = supabase_payload["cbenef_rules"]

    taxonomy = build_taxonomy(active_groups, group_keywords)
    rules_by_exact, rules_by_prefix = build_engine_tables(rules)

    results: list[dict[str, Any]] = []
    status_counts = {"BATEU": 0, "DIVERGIU": 0, "INCONCLUSIVO": 0}
    cbenef_expected_count = 0
    cbenef_ignored_count = 0

    for row in rows:
        response = classify_row(row, rules_by_exact, rules_by_prefix, taxonomy)
        conflict_reason = conflict_map.get(build_validation_key(row), {}).get("reason")
        status, reason = compare_row(row, response, conflict_reason)

        expected_cbenef = normalize_cbenef(row["cBenef"])
        if expected_cbenef:
            cbenef_expected_count += 1
        else:
            cbenef_ignored_count += 1

        result_row = {
            "Código": stringify_code(row["Código"]),
            "Descrição": clean_text(row["Descrição"]),
            "NCM_esperado": normalize_ncm(row["NCM"]),
            "NCM_retornado": normalize_ncm(response.get("matched_ncm")),
            "CST_esperado": normalize_cst(row["CST_ICMS<S>"]),
            "CST_retornado": normalize_cst(response.get("final_cst_icms")),
            "%ICMS_esperado": normalize_percent(row["% ICMS<S>"]),
            "%ICMS_retornado": normalize_percent(response.get("output_icms_rate")),
            "TRIB_esperado": normalize_trib(row["Trib."]),
            "TRIB_retornado": normalize_trib(response.get("output_trib_code")),
            "CFOP_esperado": normalize_cfop(row["CFOP"]),
            "CFOP_retornado": normalize_cfop(response.get("output_cfop")),
            "cBenef_esperado": expected_cbenef,
            "cBenef_retornado": normalize_cbenef(response.get("cbenef_code")),
            "status_validacao": status,
            "motivo_divergencia": reason,
            "CEST": normalize_digits(row["CEST"]),
            "Dep.Nome": clean_text(row["Dep.Nome"]),
            "Tipo produto": clean_text(row["Tipo produto"]),
            "_source_row": row["_source_row"],
            "_validation_key": build_validation_key(row),
        }
        results.append(result_row)
        status_counts[status] += 1

    output_payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_file": str(args.input.resolve()),
        "source_sheet": args.sheet,
        "seam": (
            "Local validator that mirrors the current get-cbenef scoring path "
            "with Supabase anon-table reads for cbenef_rules, classification_groups, "
            "and classification_group_keywords."
        ),
        "counts": {
            "total": len(results),
            "bateu": status_counts["BATEU"],
            "divergiu": status_counts["DIVERGIU"],
            "inconclusivo": status_counts["INCONCLUSIVO"],
            "cbenef_com_espera": cbenef_expected_count,
            "cbenef_sem_comparacao": cbenef_ignored_count,
            "grupos_em_conflito": len(conflict_map),
        },
        "conflicts": conflict_map,
        "columns": OUTPUT_COLUMNS,
        "rows": results,
    }

    json_path = args.output_dir / "validation_results.json"
    json_path.write_text(json.dumps(output_payload, ensure_ascii=False, indent=2), encoding="utf-8")

    summary_path = args.output_dir / "validation_summary.json"
    summary_path.write_text(
        json.dumps(
            {
                "generated_at": output_payload["generated_at"],
                "source_file": output_payload["source_file"],
                "counts": output_payload["counts"],
                "seam": output_payload["seam"],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    builder = args.builder.resolve()
    output_xlsx = args.output_dir / "cbenef_validation_report.xlsx"
    subprocess.run(
        [
            "node",
            str(builder),
            "--input-json",
            str(json_path),
            "--output-xlsx",
            str(output_xlsx),
        ],
        check=True,
    )

    print(json.dumps(output_payload["counts"], ensure_ascii=False))
    print(str(output_xlsx.resolve()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
